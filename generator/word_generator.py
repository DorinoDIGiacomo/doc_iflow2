import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from generator.sections.cover import add_cover
from generator.sections.header import add_header
from generator.sections.footer import add_footer
from generator.sections.document_control import add_document_control
from generator.sections.table_of_contents import add_table_of_contents
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ==========================================================
# Helper Functions
# ==========================================================

def add_title(document, text):

    # Space before every major section
    document.add_paragraph()

    # Top divider
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(1)

    run = p.add_run("─" * 70)
    run.font.name = "Cambria"
    run.font.size = Pt(10)
    run.font.color.rgb = __import__("docx.shared").shared.RGBColor(150, 150, 150)

    # Heading
    heading = document.add_heading(level=1)

    heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    run = heading.add_run(text)

    run.font.name = "Cambria"
    run.font.size = Pt(18)
    run.bold = True

    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(1)

    # Bottom divider
    p = document.add_paragraph()

    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)

    run = p.add_run("─" * 70)
    run.font.name = "Cambria"
    run.font.size = Pt(10)
    run.font.color.rgb = __import__("docx.shared").shared.RGBColor(150, 150, 150)

def add_subtitle(document, text):

    heading = document.add_heading(level=2)

    run = heading.add_run(text)

    run.bold = True
    run.font.name = "Cambria"
    run.font.size = Pt(14)

    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)

def add_text(document, text):

    if not text:
        text = "Not Available"

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.2

    run = paragraph.add_run(text)

    run.font.name = "Aptos"
    run.font.size = Pt(11)

def add_bullet_list(document, items):

    if not items:
        document.add_paragraph("Not Available")
        return

    for item in items:

        p = document.add_paragraph(style="List Bullet")

        p.paragraph_format.space_after = Pt(3)

        run = p.add_run(item)

        run.font.name = "Aptos"
        run.font.size = Pt(11)

def add_process_steps(document, steps):

    if not steps:

        document.add_paragraph("Not Available")
        return

    for index, step in enumerate(steps, start=1):

        title = step.get(
            "title",
            f"Step {index}"
        )

        description = step.get(
            "description",
            "Not Available"
        )

        # Step Heading
        heading = document.add_heading(
            level=2
        )

        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.space_after = Pt(2)

        run = heading.add_run(
            f"Step {index} - {title}"
        )

        run.bold = True
        run.font.name = "Cambria"
        run.font.size = Pt(13)

        # Step Description
        paragraph = document.add_paragraph()

        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15

        run = paragraph.add_run(description)

        run.font.name = "Aptos"
        run.font.size = Pt(11)

def add_numbered_process(document, steps):

    if not steps:
        return

    for index, step in enumerate(steps, start=1):

        title = step.get(
            "title",
            f"Step {index}"
        )

        description = step.get(
            "description",
            ""
        )

        document.add_heading(
            f"Step {index}: {title}",
            level=2
        )

        document.add_paragraph(
            description
        )


def add_process_flow_image(document):

    image_path = "output/process_flow.png"

    if os.path.exists(image_path):

        document.add_picture(
            image_path,
            width=Inches(6)
        )

from docx.shared import RGBColor

def add_interface_summary(document, interface):

    add_title(document, "Interface Summary")

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    header = table.rows[0].cells
    header[0].text = "Property"
    header[1].text = "Value"

    data = [
    ("Interface Name", interface.name),
    ("Description", interface.description),
    ("Runtime Platform", interface.runtime),
    ("Package", interface.package_name),
    ("Artifact ID", interface.artifact_id),
    ("Namespace", interface.namespace),

    ("Sender System(s)",
        ", ".join(interface.sender_systems)
        if interface.sender_systems else "Not Available"),

    ("Receiver System(s)",
        ", ".join(interface.receiver_systems)
        if interface.receiver_systems else "Not Available"),

    ("External System(s)",
        ", ".join(interface.external_systems)
        if interface.external_systems else "Not Available"),

    ("Endpoints",
        ", ".join(interface.endpoints)
        if interface.endpoints else "Not Available"),

    ("Sender Adapter(s)",
        ", ".join(interface.sender_adapters)
        if interface.sender_adapters else "Not Available"),

    ("Receiver Adapter(s)",
        ", ".join(interface.receiver_adapters)
        if interface.receiver_adapters else "Not Available"),

    ("Communication Protocol",
        ", ".join(interface.protocols)
        if interface.protocols else "Not Available"),

    ("Message Format",
        ", ".join(interface.message_formats)
        if interface.message_formats else "Not Available"),

    ("Authentication",
        ", ".join(interface.authentication)
        if interface.authentication else "Not Available")
]

    for key, value in data:

        if not value or str(value).strip() == "" or value == "Not Available":
            continue

        row = table.add_row().cells
        row[0].text = key
        row[1].text = str(value)

    document.add_page_break()
    
def add_system_landscape(document, interface):

    add_title(document, "System Landscape")

    source = ", ".join(interface.sender_systems) if interface.sender_systems else "Source System"
    target = ", ".join(interface.receiver_systems) if interface.receiver_systems else "Target System"

    def draw_box(text):
        table = document.add_table(rows=1, cols=1)
        table.style = "Table Grid"

        cell = table.cell(0, 0)
        p = cell.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        run = p.add_run(text)
        run.bold = True
        run.font.name = "Cambria"
        run.font.size = Pt(11)

    # Source
    draw_box(f"Source System\n{source}")

    p = document.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.add_run("│")
    p.add_run("\n▼")

    # CPI
    draw_box("SAP Integration Suite\nCloud Integration")

    p = document.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.add_run("│")
    p.add_run("\n▼")

    # Receiver
    draw_box(f"Receiver System\n{target}")

    document.add_paragraph()   

# ==========================================================
# AI Generated Sections
# ==========================================================

def add_ai_sections(document, ai_data):

    if not ai_data:
        return

    # ======================================================
    # 1. Introduction
    # ======================================================

    add_title(document, "1. Introduction")

    add_subtitle(document, "1.1 Purpose")

    add_text(
        document,
        ai_data.get("purpose", "")
    )

    add_subtitle(document, "1.2 Scope")

    add_text(
        document,
        ai_data.get("scope", "")
    )

    # ======================================================
    # 2. Business Context
    # ======================================================

    add_title(document, "2. Business Context")

    add_text(
        document,
        ai_data.get("business_context", "")
    )

    # ======================================================
    # 3. Business Challenges
    # ======================================================

    add_title(document, "3. Business Challenges")

    add_bullet_list(
        document,
        ai_data.get(
            "business_challenges",
            []
        )
    )

    # ======================================================
    # 4. Proposed Solution
    # ======================================================

    add_title(document, "4. Proposed Solution")

    add_subtitle(
        document,
        "4.1 Solution Summary"
    )
    
    add_text(
        document,
        ai_data.get(
            "solution_summary",
            ""
        )
    )
    
    add_subtitle(
        document,
        "4.2 Process Overview"
    )

    add_text(
        document,
        ai_data.get(
            "process_overview",
            ""
        )
    )

    # ======================================================
    # 5. Solution Architecture
    # ======================================================

    add_title(document, "5. Solution Architecture")

    add_text(
        document,
        ai_data.get(
            "architecture_overview",
            ""
        )
    )
    
    document.add_page_break()

    # ======================================================
    # 6. Process Flow Diagram
    # ======================================================
    
    add_process_flow_image(
    document
    )
    
    document.add_page_break()

    # ======================================================
    # 7. End-to-End Integration Flow
    # ======================================================

    add_title(
        document,
        "7. End-to-End Integration Flow"
    )

    add_process_steps(
        document,
        ai_data.get(
        "process_flow",
        []
        )
    )

    # ======================================================
    # 8. Security Design
    # ======================================================
    
    add_title(
        document,
        "8. Security Design"
    )

    add_text(
        document,
        ai_data.get(
            "security_design",
            ""
        )
    )

    # ======================================================
    # 9. Monitoring & Error Handling
    # ======================================================

    add_title(
        document,
        "9. Monitoring & Error Handling"
    )

    add_subtitle(
        document,
        "9.1 Monitoring"
    )

    add_text(
        document,
        ai_data.get(
            "monitoring",
            ""
        )
    )

    add_subtitle(
        document,
        "9.2 Error Handling"
    )

    add_text(
        document,
        ai_data.get(
            "error_handling",
            ""
        )
    )

    # ======================================================
    # 10. Business Benefits
    # ======================================================

    add_title(
        document,
        "10. Business Benefits"
    )

    add_bullet_list(
        document,
        ai_data.get(
            "business_benefits",
            []
        )
    )

    # ======================================================
    # 11. Assumptions
    # ======================================================

    add_title(
        document,
        "11. Assumptions"
    )

    add_bullet_list(
        document,
        ai_data.get(
            "assumptions",
            []
        )
    )

    # ======================================================
    # 12. Limitations
    # ======================================================

    add_title(
        document,
        "12. Limitations"
    )

    add_bullet_list(
        document,
        ai_data.get(
            "limitations",
            []
        )
    )

    # ======================================================
    # 13. Recommendations
    # ======================================================

    add_title(
        document,
        "13. Recommendations"
    )

    add_bullet_list(
        document,
        ai_data.get(
            "recommendations",
            []
        )
    )

    # ======================================================
    # 14. Conclusion
    # ======================================================

    add_title(
        document,
        "14. Conclusion"
    )

    add_text(
        document,
        ai_data.get(
            "conclusion",
            ""
        )
    )

# ==========================================================
# Auto Update Fields
# ==========================================================
    
def enable_auto_update_fields(document):
    """
    Instruct Microsoft Word to update all fields (TOC, page numbers, etc.)
    when the document is opened.
    """

    settings = document.settings.element

    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")

    settings.append(update_fields)
 
# ==========================================================
# Main Word Generator
# ==========================================================

def generate_word(interface, output_path, ai_data=None):

    print("\n========== Word Generator ==========\n")

    document = Document()

    from docx.shared import Inches

    section = document.sections[0]

    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    # ======================================================
    # Header
    # ======================================================

    add_header(document)

    # ======================================================
    # Footer
    # ======================================================    

    add_footer(document)

    # ======================================================
    # Cover Page
    # ======================================================

    print("Generating Cover Page...")

    add_cover(
        document,
        interface
    )

    print("Generating Document Control...")

    add_document_control(
        document,
        interface
    )
    
    print("Generating Table of Contents...")
    
    add_table_of_contents(
        document
    )

    print("Generating Interface Summary...")

    add_interface_summary(
        document,
        interface
    )

    #print("Generating System Landscape...")

    #add_system_landscape(
        #document,
        #interface
    #)

# ======================================================
# AI Generated Sections
# ======================================================

    print("Generating AI Documentation...")

    add_ai_sections(
        document,
        ai_data
    )
    
# ======================================================
# Save Document
# ======================================================

    print("Saving Solution Design Document...")

    # Enable automatic field updates (TOC, page numbers, etc.)
    enable_auto_update_fields(document)

    os.makedirs(
        os.path.dirname(output_path),
        exist_ok=True
    )

    document.save(
        output_path
    )

    print("\n" + "=" * 60)
    print(" SAP CPI Solution Design Document Generated Successfully")
    print("=" * 60)
    print(f"Output File : {output_path}")
    print("=" * 60)