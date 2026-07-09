from docx.shared import Pt


def add_cover(document, interface):

    title = document.add_heading(
        "SAP CPI Technical Specification",
        level=0
    )

    title.runs[0].font.size = Pt(22)

    document.add_paragraph()

    document.add_heading(interface.name, level=1)

    document.add_paragraph(f"Version : {interface.version}")

    document.add_page_break()


def add_interface_summary(document, interface):

    document.add_heading("1. Interface Summary", level=1)

    table = document.add_table(rows=3, cols=2)

    table.style = "Table Grid"

    table.cell(0,0).text = "Name"
    table.cell(0,1).text = interface.name

    table.cell(1,0).text = "Version"
    table.cell(1,1).text = interface.version

    table.cell(2,0).text = "Description"
    table.cell(2,1).text = interface.description


def add_component_table(document, interface):

    document.add_heading("2. Components", level=1)

    table = document.add_table(rows=1, cols=4)

    table.style = "Table Grid"

    hdr = table.rows[0].cells

    hdr[0].text = "Name"
    hdr[1].text = "Category"
    hdr[2].text = "BPMN"
    hdr[3].text = "Process"

    for component in interface.components:

        row = table.add_row().cells

        row[0].text = component.name
        row[1].text = component.category
        row[2].text = component.bpmn_type
        row[3].text = component.process