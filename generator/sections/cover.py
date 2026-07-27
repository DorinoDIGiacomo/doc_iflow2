from datetime import datetime
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt


def add_cover(document, interface):

    # =====================================================
    # Optional Logo (Top Left)
    # =====================================================

    logo_path = Path("assets/logo.png")

    if logo_path.exists():
        logo = document.add_paragraph()
        logo.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        logo.add_run().add_picture(str(logo_path), width=Inches(1.2))

    # =====================================================
    # Main Title
    # =====================================================

    document.add_paragraph()

    title = document.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = title.add_run("SAP Integration Suite")
    run.bold = True
    run.font.name = "Cambria"
    run.font.size = Pt(26)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = subtitle.add_run("Interface Design Document")
    run.bold = True
    run.font.name = "Cambria"
    run.font.size = Pt(18)

    document.add_paragraph()

    # =====================================================
    # Interface Name
    # =====================================================

    interface_name = document.add_paragraph()
    interface_name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    run = interface_name.add_run(interface.name)
    run.bold = True
    run.font.name = "Cambria"
    run.font.size = Pt(16)

    document.add_paragraph()

    # =====================================================
    # Information Table
    # =====================================================

    table = document.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    rows = [
        ("Interface Name", interface.name),
        ("Platform", "SAP Integration Suite"),
        ("Runtime", "Cloud Integration"),
        ("Generated On", datetime.now().strftime("%d-%b-%Y")),
        ("Document Status", "Draft"),
    ]

    for row, (label, value) in zip(table.rows, rows):

        left = row.cells[0].paragraphs[0]
        left.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        run = left.add_run(label)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(11)

        right = row.cells[1].paragraphs[0]
        right.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        run = right.add_run(value)
        run.font.name = "Aptos"
        run.font.size = Pt(11)

    document.add_paragraph()

    # =====================================================
    # Description
    # =====================================================

    description = document.add_paragraph()
    description.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

    run = description.add_run(
        "This document provides the functional and technical design of the "
        "SAP Integration Suite interface. It includes the integration "
        "architecture, process flow, component configuration, and "
        "implementation details."
    )

    run.font.name = "Aptos"
    run.font.size = Pt(11)

    # =====================================================
    # Next Page
    # =====================================================

    document.add_page_break()