from datetime import datetime
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def add_document_control(document, interface):

    document.add_heading(
        "Document Control",
        level=1
    )

    # =====================================================
    # Document Information
    # =====================================================

    document.add_heading(
        "Document Information",
        level=2
    )

    table = document.add_table(
        rows=13,
        cols=2
    )

    table.style = "Table Grid"

    data = [

        ("Document Title", "SAP Integration Suite Interface Design Document"),

        ("Interface Name", interface.name),

        ("Version", interface.version or "1.0"),

        ("Description", interface.description or "Not Available"),

        ("Prepared By", "AI SAP CPI Documentation Generator"),

        ("Generated On", datetime.now().strftime("%d-%m-%Y %H:%M")),

        ("Platform", "SAP Integration Suite (Cloud Integration)"),
        
        ("Runtime", interface.runtime or "Cloud Integration"),
        
        ("Document Status", "Draft"),
        
        ("Classification", "Internal"),
        
        ("Package", interface.package_name or "Not Available"),
        
        ("Artifact ID", interface.artifact_id or "Not Available"),
        
        ("Namespace", interface.namespace or "Not Available")

    ]

    for row, (key, value) in enumerate(data):

        table.cell(row, 0).text = key

        table.cell(row, 1).text = str(value)

    document.add_paragraph()

    # =====================================================
    # Revision History
    # =====================================================

    document.add_heading(
        "Revision History",
        level=2
    )

    revision = document.add_table(
        rows=2,
        cols=5
    )

    revision.style = "Table Grid"

    headers = [
        "Version",
        "Date",
        "Author",
        "Description",
        "Status"
    ]

    for col, header in enumerate(headers):

        revision.cell(0, col).text = header

    revision.cell(1, 0).text = interface.version or "1.0"

    revision.cell(1, 1).text = datetime.now().strftime("%d-%m-%Y")

    revision.cell(1, 2).text = "SAP CPI Interface Design Generator"

    revision.cell(1, 3).text = "Initial Interface Design Document"

    revision.cell(1, 4).text = "Draft"

    document.add_paragraph()

    # =====================================================
    # Approvals
    # =====================================================

    document.add_heading(
        "Approvals",
        level=2
    )

    approval = document.add_table(
        rows=2,
        cols=3
    )

    approval.style = "Table Grid"

    approval.cell(0, 0).text = "Role"
    approval.cell(0, 1).text = "Name"
    approval.cell(0, 2).text = "Signature"

    approval.cell(1, 0).text = "Project / Integration Lead"
    approval.cell(1, 1).text = ""
    approval.cell(1, 2).text = ""

    document.add_page_break()