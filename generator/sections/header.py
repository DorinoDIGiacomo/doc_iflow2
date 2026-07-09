from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def add_header(document):

    section = document.sections[0]

    # Different header/footer for cover page
    section.different_first_page_header_footer = True

    # Add logo to BOTH the cover page header and the remaining pages
    for header in [section.first_page_header, section.header]:

        paragraph = header.paragraphs[0]

        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        paragraph.add_run().add_picture(
            "generator/assets/truspeq_logo.png",
            width=Inches(1.8)
        )