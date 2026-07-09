from docx.shared import Pt
from utils.property_table_parser import parse_property_table


def add_component_details(document, interface):

    document.add_heading(
        "Component Details",
        level=1
    )

    for index, component in enumerate(interface.components, start=1):

        # =====================================================
        # Component Heading
        # =====================================================

        document.add_heading(
            f"{index}. {component.name}",
            level=2
        )

        # =====================================================
        # Basic Information
        # =====================================================

        info = document.add_table(
            rows=0,
            cols=2
        )

        info.style = "Table Grid"

        def add_row(key, value):

            row = info.add_row().cells

            row[0].text = str(key)

            row[1].text = str(value)

        analysis = component.analysis or {}

        add_row(
            "Category",
            component.category
        )

        add_row(
            "Process",
            component.process
        )

        add_row(
            "Direction",
            analysis.get(
                "direction",
                "-"
            )
        )

        add_row(
            "Purpose",
            analysis.get(
                "purpose",
                "-"
            )
        )

        add_row(
            "Incoming",
            ", ".join(
                analysis.get(
                    "incoming",
                    []
                )
            ) or "-"
        )

        add_row(
            "Outgoing",
            ", ".join(
                analysis.get(
                    "outgoing",
                    []
                )
            ) or "-"
        )

        # =====================================================
        # Properties
        # =====================================================

        if component.properties:

            document.add_paragraph()

            heading = document.add_paragraph()

            run = heading.add_run("Properties")

            run.bold = True

            run.font.size = Pt(11)

            property_table = document.add_table(
                rows=1,
                cols=2
            )

            property_table.style = "Table Grid"

            header = property_table.rows[0].cells

            header[0].text = "Property"

            header[1].text = "Value"

            for key, value in component.properties.items():

                parsed = parse_property_table(
                    str(value)
                )

                if parsed:

                    row = property_table.add_row().cells

                    row[0].text = key

                    row[1].text = "Structured Table"

                    headers = list(parsed[0].keys())

                    nested = document.add_table(
                        rows=1,
                        cols=len(headers)
                    )

                    nested.style = "Table Grid"

                    hdr = nested.rows[0].cells

                    for i, column in enumerate(headers):

                        hdr[i].text = column

                    for item in parsed:

                        cells = nested.add_row().cells

                        for i, column in enumerate(headers):

                            cells[i].text = str(
                                item.get(
                                    column,
                                    ""
                                )
                            )

                else:

                    row = property_table.add_row().cells

                    row[0].text = key

                    row[1].text = str(value)

        document.add_page_break()